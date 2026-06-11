import json
import os
import base64
import io
import re
import datetime
from pathlib import Path
from flask import Flask, render_template, request, jsonify, Response
from flask_sqlalchemy import SQLAlchemy
from dotenv import load_dotenv
import anthropic
import openpyxl
import pdfplumber
from PIL import Image
import fitz  # pymupdf

load_dotenv(override=True)

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 50 * 1024 * 1024  # 50MB

BASE_DIR = Path(__file__).parent
PRICES_FILE = BASE_DIR / "prices.json"
CONTRACTORS_FILE = BASE_DIR / "contractors.json"

# SQLite(ローカル) または PostgreSQL(本番) を DATABASE_URL で切り替え
database_url = os.environ.get("DATABASE_URL", f"sqlite:///{BASE_DIR}/app.db")
# Render の PostgreSQL URL は "postgres://" で始まるため修正が必要
if database_url.startswith("postgres://"):
    database_url = database_url.replace("postgres://", "postgresql://", 1)

app.config["SQLALCHEMY_DATABASE_URI"] = database_url
app.config["SQLALCHEMY_TRACK_MODIFICATIONS"] = False
db = SQLAlchemy(app)


class Price(db.Model):
    __tablename__ = "prices"
    id       = db.Column(db.Integer, primary_key=True)
    category = db.Column(db.String(100), nullable=False)
    item     = db.Column(db.String(200), nullable=False)
    unit     = db.Column(db.String(20), default="式")
    price    = db.Column(db.Integer, nullable=False)
    memo     = db.Column(db.String(500), default="")

    def to_dict(self):
        return {
            "id": self.id,
            "category": self.category,
            "item": self.item,
            "unit": self.unit,
            "price": self.price,
            "memo": self.memo or ""
        }


class Contractor(db.Model):
    __tablename__ = "contractors"
    id                 = db.Column(db.Integer, primary_key=True)
    name               = db.Column(db.String(200), nullable=False, unique=True)
    categories         = db.Column(db.JSON, default=list)
    estimate_count     = db.Column(db.Integer, default=0)
    last_estimate_date = db.Column(db.String(20), default="")
    rating             = db.Column(db.Integer, default=0)
    memo               = db.Column(db.String(500), default="")

    def to_dict(self):
        return {
            "id": self.id,
            "name": self.name,
            "categories": self.categories or [],
            "estimate_count": self.estimate_count,
            "last_estimate_date": self.last_estimate_date or "",
            "rating": self.rating,
            "memo": self.memo or ""
        }


class AnalysisHistory(db.Model):
    __tablename__ = "analysis_history"
    id               = db.Column(db.Integer, primary_key=True)
    analyzed_at      = db.Column(db.String(30), nullable=False)
    contractor_name  = db.Column(db.String(200), default="")
    file_names       = db.Column(db.String(500), default="")
    missing_count    = db.Column(db.Integer, default=0)
    overpriced_count = db.Column(db.Integer, default=0)
    summary          = db.Column(db.Text, default="")
    result_json      = db.Column(db.Text, default="")

    def to_dict(self):
        return {
            "id": self.id,
            "analyzed_at": self.analyzed_at,
            "contractor_name": self.contractor_name or "",
            "file_names": self.file_names or "",
            "missing_count": self.missing_count,
            "overpriced_count": self.overpriced_count,
            "summary": self.summary or ""
        }


with app.app_context():
    db.create_all()
    # テーブルが空の場合のみ既存JSONデータを移行
    if Price.query.count() == 0 and PRICES_FILE.exists():
        with open(PRICES_FILE, encoding="utf-8") as f:
            for p in json.load(f):
                db.session.add(Price(
                    id=p["id"],
                    category=p["category"],
                    item=p["item"],
                    unit=p.get("unit", "式"),
                    price=p["price"],
                    memo=p.get("memo", "")
                ))
        db.session.commit()
    if Contractor.query.count() == 0 and CONTRACTORS_FILE.exists():
        with open(CONTRACTORS_FILE, encoding="utf-8") as f:
            for c in json.load(f):
                db.session.add(Contractor(
                    id=c["id"],
                    name=c["name"],
                    categories=c.get("categories", []),
                    estimate_count=c.get("estimate_count", 0),
                    last_estimate_date=c.get("last_estimate_date", "") or "",
                    rating=c.get("rating", 0),
                    memo=c.get("memo", "")
                ))
        db.session.commit()


@app.before_request
def check_auth():
    username = os.environ.get("APP_USERNAME")
    password = os.environ.get("APP_PASSWORD")
    if not username or not password:
        return
    auth = request.authorization
    if not auth or auth.username != username or auth.password != password:
        return Response(
            "認証が必要です", 401,
            {"WWW-Authenticate": 'Basic realm="Estimate Checker"'}
        )


SAMPLE_EMAILS = """
--- メール例① ---
株式会社カイセイプランニング
田村　様

お世話になっております。
コスモバンク三木でございます。

表題の件について
下記の通りご訂正いただくことは可能でしょうか。

【不要】
・洋室　面格子目隠し　材工
【追加】
・各窓なければカーテンレール新設
【変更】
・エアコンクリーニング　⇒　エアコン交換（材工）

以上になります。
お忙しいところ恐縮ですが
ご確認のほどよろしくお願い申し上げます。

三木

--- メール例② ---
有限会社マルミヤ商会

宮島　様

お世話になっております。
コスモバンク三木でございます。

表題の件について
下記の通りご訂正いただくことは可能でしょうか。

【追加】
・キッチン換気扇交換
・長押の付着物処分
・冷蔵庫用コンセント新設
※モニターホン含め参考設置場所写真添付いたします。
【変更】
・エアコンクリーニング　⇒　エアコン交換（材工）

以上になります。
お忙しいところ恐縮ですが
ご確認のほどよろしくお願い申し上げます。

三木

--- メール例③ ---
株式会社ヒロ・コーポレーション

小野　様

お世話になっております。
コスモバンク三木でございます。
先程はお電話にてありがとうございました。

表題の件について
下記の通りご訂正いただくことは可能でしょうか。
単価に関してはあくまでも希望ですのでご承知おきください。

【単価調整希望】
・ボールタップ交換　20,000　⇒　15,000

また、別紙にて浴室ドアの交換のお見積りもいただきたく存じます。

以上になります。
お忙しいところ恐縮ですが
ご確認のほどよろしくお願い申し上げます。

三木

--- メール例④ ---
合同会社杉田商店
ご担当者　様

お世話になっております。
コスモバンク三木でございます。

表題の件について、下記の内容をご確認いただけますでしょうか。

【洋室】
・フロアタイル敷設　⇒　CF上貼りまたは張替え
【キッチン】
・フロアタイル敷設　⇒　CF上貼りまたは張替え

以上になります。
お忙しいところ恐縮ですが
ご確認のほどよろしくお願い申し上げます。

三木

--- メール例⑤ ---
株式会社TOP FLASH
山口　様

いつも大変お世話になっております。
コスモバンク株式会社の三木と申します。

表題の件について
ウォシュレット新設（材工）を
28,000円希望で追加いただきたく存じます。

お忙しいところ恐縮ですが
ご確認のほどよろしくお願い申し上げます。

三木
"""

MANUAL_TEXT = """
【水漏れ確認】
・水栓＝ハンドル、吐水根本、シャワーヘッド、シャワー根本、止水栓、トイレタンク内・便器内
・排水＝排水トラップ（キッチン、洗面台、ユニットバス洗面器 ※主に浴槽との結合部、トイレ洗浄管）
・浴槽および洗面台の排水ゴム栓漏れ確認
■見積最初に、5～10分は流水して水栓、排水の漏れを確認。退出前に再度確認する。
■見落としがちだが給湯器も開栓して給湯ハンドル側も確認する。

【建付け確認】定義は開閉のスムーズさ、異音や閉めた時の隙間
・扉（トイレ、浴室含む）＝開閉不良、枠接触、ラッチかかり具合、取ってグラつき、カギ施錠不良
・引戸＝スライド不良、閉めた際の隙間、取ってグラつき
・折戸（浴室含む）＝開閉不良、吊元グラつき・脱線確認、取手グラつき
・蝶番＝開閉不良
・網戸＝スライド不良、戸車確認、閉めた際の隙間、強度、落下防止部材確認
・サッシ＝クレセント不良、スライド不良、戸車確認、鍵確認
・シャッター雨戸＝開閉不良、シャッター錠不良、引き紐確認
・スライド雨戸＝スライド不良、閉めた際の隙間、戸車確認、鍵確認

【設備品】※要動作確認
・エアコン＝なければ1台は必ず設置。2台以上であれば撤去も検討。専用回路になっているかと製造年も要確認。
  ■エアコンに関しては製造日から10年過ぎていれば交換提案をする。
・天井シーリングライト（LED）＝メインの部屋に1台設置。既存にリモコンが無い、LEDではないかも要確認。
・火災報知器＝キッチン（熱）・室内（煙）が各室設置されているか確認。製造日も要確認。
  ※自火報が付いている場合は不要。
・モニターホン＝通話確認、チャイム音確認、モニター画像確認
  玄関チャイム、ドアチャイム、インターホン（モニター無し）はモニターホン提案。
  ※既存のチャイムの撤去と下地処理（プレート処理等）は必須
・便座＝温水洗浄便座提案、割れ確認、便座ゴム確認、温水洗浄便座の動作確認
・換気扇＝吸い込み確認、異音確認
・IH＝動作確認（鍋などで確認）
・リモコン類＝液晶・動作確認
・スイッチ＝動作確認、感触不具合
・照明＝玉切れ、点灯確認
・洗濯機水栓＝オートストップニップルか確認

【表層】
・クロス、CF、襖、畳、巾木など貼替が必要な箇所はご提案ください。
  （クリーニング、簡易補修での対応で問題ない場合はそちらでご提案ください）
・建具表層の破損等確認。
・和室の洋間変更の提案

【その他】
・床鳴り＝床鳴りがあれば必ず報告・提案。
・異臭＝排水からの異臭がある場合臭気ゴムなど要確認
・コーキング＝キッチン・浴室他 カビ・汚れ確認
・排水ゴム栓＝劣化確認、サイズ感があっているか確認
・防水パン＝Lボ 又はLボバンドがあるか確認
・TVジャック 通信不具合確認（アンテナチェッカー仕様にて）
・換気扇及びレンジフード排気口のベントキャップが防火ダンパー付きの場合
  温度ヒューズが切れていないか確認。また接続はずれやダクト破れの有無の確認

●確認ポイント●
・入居者が生活するうえで支障が出るかを判断基準として各所修繕必要かの判断を行う。
  ※けがの心配がある、使えないなど
・工事に伴い仕上がりにも気を使うようお願いいたします。
  ※露出配線など配線のむき出し、垂れ流しは控えモールなどで見た目よく処理する。
"""


def auto_register_contractor(name: str, categories: list):
    """見積書から検出した業者名を自動登録。既存なら estimate_count と categories を更新。"""
    if not name:
        return
    existing = Contractor.query.filter_by(name=name).first()
    today = datetime.date.today().isoformat()
    if existing:
        existing.estimate_count += 1
        existing.last_estimate_date = today
        cats = list(existing.categories or [])
        for cat in categories:
            if cat not in cats:
                cats.append(cat)
        existing.categories = cats
    else:
        new_id = (db.session.query(db.func.max(Contractor.id)).scalar() or 0) + 1
        db.session.add(Contractor(
            id=new_id,
            name=name,
            categories=list(set(categories)),
            estimate_count=1,
            last_estimate_date=today,
            rating=0,
            memo=""
        ))
    db.session.commit()


def extract_excel_text(file_bytes):
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    lines = []
    for sheet in wb.worksheets:
        lines.append(f"=== シート: {sheet.title} ===")
        for row in sheet.iter_rows():
            cells = [str(cell.value) for cell in row if cell.value is not None]
            if cells:
                lines.append("\t".join(cells))
    return "\n".join(lines)


def extract_pdf_text(file_bytes):
    text_parts = []
    with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
        for page in pdf.pages:
            text = page.extract_text()
            if text:
                text_parts.append(text)
    return "\n".join(text_parts)


def image_to_base64(file_bytes, media_type):
    img = Image.open(io.BytesIO(file_bytes))
    # 大きすぎる画像はリサイズ（Claude APIの制限対策）
    max_size = 1568
    if max(img.width, img.height) > max_size:
        img.thumbnail((max_size, max_size), Image.LANCZOS)
    buf = io.BytesIO()
    fmt = "JPEG" if media_type in ("image/jpeg", "image/jpg") else "PNG"
    img.save(buf, format=fmt)
    return base64.standard_b64encode(buf.getvalue()).decode("utf-8")


def pdf_pages_to_images(file_bytes):
    """PDFの各ページをJPEG画像に変換してbase64リストで返す"""
    doc = fitz.open(stream=file_bytes, filetype="pdf")
    images = []
    for page in doc:
        mat = fitz.Matrix(2, 2)  # 2倍ズームで解像度確保
        pix = page.get_pixmap(matrix=mat)
        img_bytes = pix.tobytes("jpeg")
        b64 = base64.standard_b64encode(img_bytes).decode("utf-8")
        images.append(b64)
    return images


def build_price_master_text(prices):
    if not prices:
        return "（単価マスタ未登録）"
    lines = ["【社内単価マスタ】"]
    current_category = None
    for p in sorted(prices, key=lambda x: x["category"]):
        if p["category"] != current_category:
            current_category = p["category"]
            lines.append(f"\n[{current_category}]")
        lines.append(f"  {p['item']}: {p['price']:,}円/{p['unit']}")
    return "\n".join(lines)


@app.route("/")
def index():
    return render_template("index.html")


@app.route("/analyze", methods=["POST"])
def analyze():
    api_key = os.environ.get("ANTHROPIC_API_KEY")
    if not api_key:
        return jsonify({"error": "ANTHROPIC_API_KEY が設定されていません。.env ファイルを確認してください。"}), 500

    estimate_files = request.files.getlist("estimates")
    photo_files = request.files.getlist("photos")

    if not estimate_files or all(f.filename == "" for f in estimate_files):
        return jsonify({"error": "見積書ファイルをアップロードしてください。"}), 400

    estimate_texts = []
    for f in estimate_files:
        if f.filename == "":
            continue
        file_bytes = f.read()
        filename = f.filename.lower()
        try:
            if filename.endswith((".xlsx", ".xls")):
                text = extract_excel_text(file_bytes)
            elif filename.endswith(".pdf"):
                text = extract_pdf_text(file_bytes)
            else:
                return jsonify({"error": f"対応していないファイル形式です: {f.filename}"}), 400
        except Exception as e:
            return jsonify({"error": f"ファイルの読み込みに失敗しました（{f.filename}）: {e}"}), 400
        estimate_texts.append(f"【ファイル名: {f.filename}】\n{text}")

    combined_estimate_text = "\n\n".join(estimate_texts)

    prices = [p.to_dict() for p in Price.query.order_by(Price.id).all()]
    price_master_text = build_price_master_text(prices)

    system_prompt = f"""あなたは不動産リノベーション工事の見積精査の専門家です。
以下の「見積り内容マニュアル」と「社内単価マスタ」に基づいて、業者からの見積書を精査してください。

{MANUAL_TEXT}

{price_master_text}

精査結果は必ず以下のJSON形式のみで返してください。説明文やマークダウンは不要です。

{{
  "missing_items": [
    {{"category": "カテゴリ名", "item": "項目名", "reason": "なぜ問題か", "action": "推奨アクション"}}
  ],
  "overpriced_items": [
    {{"item": "工事項目名", "estimate_price": 数値, "reference_price": 数値, "unit": "単位", "diff_rate": 差額率(整数%), "source": "マスタ比" または "相場比"}}
  ],
  "ok_items": [
    {{"item": "問題なしの項目名"}}
  ],
  "summary": "総評テキスト（2〜3文）",
  "estimate_prices": [
    {{"category": "カテゴリ名", "item": "見積書上の項目名", "unit": "単位", "price": 単価数値}}
  ],
  "contractor_name": "見積書の発行会社名（業者名）。見当たらない場合はnull"
}}

精査の注意点:
- missing_itemsは、マニュアルに記載があるのに見積書に完全に記載がない重要項目のみ指摘する
- overpriced_itemsは以下の2軸で判定する
  1. 社内単価マスタに項目がある場合：マスタ単価より高ければ指摘（reference_priceにマスタ単価を入れ、sourceに"マスタ比"を入れる）
  2. 社内単価マスタにない項目の場合：日本の原状回復工事・リノベーション工事の一般的な市場相場と比較して明らかに割高であれば指摘（reference_priceに相場の目安単価を入れ、sourceに"相場比"を入れる）
- 単価が明記されていない項目はoverpriced_itemsから除外する
- diff_rateは小数点なしの整数（例: 25）で入れること
- ok_itemsは見積書に記載があり問題ない項目を入れること
- estimate_pricesには見積書に単価が明記されている工事項目をすべて抽出すること。合計金額のみで単価不明の項目は含めないこと
- contractor_nameは見積書のヘッダー・社名欄・印鑑・発行者情報などから業者の会社名を抽出すること。見つからない場合はnull"""

    content_blocks = [
        {
            "type": "text",
            "text": f"以下の見積書を精査してください。\n\n【見積書内容】\n{combined_estimate_text}"
        }
    ]

    site_doc_files = request.files.getlist("site_docs")
    site_doc_images = []

    for f in site_doc_files:
        if f.filename == "":
            continue
        file_bytes = f.read()
        filename = f.filename.lower()
        if filename.endswith(".pdf"):
            try:
                pages = pdf_pages_to_images(file_bytes)
                site_doc_images.extend(pages)
            except Exception as e:
                return jsonify({"error": f"現場資料PDFの読み込みに失敗しました（{f.filename}）: {e}"}), 400
        else:
            mime = f.content_type or "image/jpeg"
            site_doc_images.append(image_to_base64(file_bytes, mime))

    if site_doc_images:
        content_blocks.append({"type": "text", "text": "以下は現場資料（室内状況報告書・間取り図等）です。"})
        for b64 in site_doc_images:
            content_blocks.append({
                "type": "image",
                "source": {"type": "base64", "media_type": "image/jpeg", "data": b64}
            })

    photo_blocks = []
    for photo in photo_files:
        if photo.filename == "":
            continue
        file_bytes = photo.read()
        mime = photo.content_type or "image/jpeg"
        b64 = image_to_base64(file_bytes, mime)
        photo_blocks.append({
            "type": "image",
            "source": {"type": "base64", "media_type": mime, "data": b64}
        })

    if photo_blocks:
        content_blocks.append({"type": "text", "text": "以下は現場写真です。"})
        content_blocks.extend(photo_blocks)

    if site_doc_images or photo_blocks:
        content_blocks.append({
            "type": "text",
            "text": "上記の現場資料・写真も考慮して精査してください。"
        })

    try:
        client = anthropic.Anthropic(api_key=api_key)
        message = client.messages.create(
            model="claude-sonnet-4-6",
            max_tokens=8192,
            system=system_prompt,
            messages=[{"role": "user", "content": content_blocks}]
        )
    except anthropic.AuthenticationError:
        return jsonify({"error": "APIキーが無効です。.env の ANTHROPIC_API_KEY を確認してください。"}), 500
    except Exception as e:
        return jsonify({"error": f"AI APIの呼び出しに失敗しました: {e}"}), 500

    raw = message.content[0].text.strip()
    json_match = re.search(r'\{.*\}', raw, re.DOTALL)
    if json_match:
        raw = json_match.group(0)

    try:
        result = json.loads(raw)
    except json.JSONDecodeError:
        return jsonify({"error": f"AIの応答を解析できませんでした。もう一度お試しください。\n\n応答内容:\n{raw[:500]}"}), 500

    contractor_name = result.get("contractor_name")
    if contractor_name:
        cats = list({p["category"] for p in result.get("estimate_prices", []) if p.get("category")})
        auto_register_contractor(contractor_name, cats)

    history = AnalysisHistory(
        analyzed_at=datetime.datetime.now().isoformat(timespec='seconds'),
        contractor_name=contractor_name or "",
        file_names=", ".join(f.filename for f in estimate_files if f.filename),
        missing_count=len(result.get("missing_items", [])),
        overpriced_count=len(result.get("overpriced_items", [])),
        summary=result.get("summary", ""),
        result_json=json.dumps(result, ensure_ascii=False)
    )
    db.session.add(history)
    db.session.commit()

    return jsonify(result)


# --- 単価マスタ API ---

@app.route("/prices")
def prices_page():
    prices = [p.to_dict() for p in Price.query.order_by(Price.id).all()]
    categories = sorted(set(p["category"] for p in prices))
    category_counts = {c: sum(1 for p in prices if p["category"] == c) for c in categories}
    return render_template("prices.html", prices=prices, categories=categories, category_counts=category_counts)


@app.route("/api/prices", methods=["GET"])
def get_prices():
    return jsonify([p.to_dict() for p in Price.query.order_by(Price.id).all()])


@app.route("/api/prices", methods=["POST"])
def add_price():
    data = request.get_json()
    new_id = (db.session.query(db.func.max(Price.id)).scalar() or 0) + 1
    entry = Price(
        id=new_id,
        category=data["category"].strip(),
        item=data["item"].strip(),
        unit=data["unit"].strip(),
        price=int(data["price"]),
        memo=data.get("memo", "").strip()
    )
    db.session.add(entry)
    db.session.commit()
    return jsonify(entry.to_dict()), 201


@app.route("/api/prices/<int:price_id>", methods=["PUT"])
def update_price(price_id):
    p = Price.query.get(price_id)
    if not p:
        return jsonify({"error": "見つかりません"}), 404
    data = request.get_json()
    p.category = data["category"].strip()
    p.item = data["item"].strip()
    p.unit = data["unit"].strip()
    p.price = int(data["price"])
    p.memo = data.get("memo", "").strip()
    db.session.commit()
    return jsonify(p.to_dict())


@app.route("/api/prices/<int:price_id>", methods=["DELETE"])
def delete_price(price_id):
    Price.query.filter_by(id=price_id).delete()
    db.session.commit()
    return jsonify({"ok": True})


@app.route("/generate-email", methods=["POST"])
def generate_email_api():
    api_key = os.environ.get("ANTHROPIC_API_KEY")
    if not api_key:
        return jsonify({"error": "ANTHROPIC_API_KEY が設定されていません。"}), 500

    data = request.get_json()
    company_name = data.get("company_name", "").strip()
    contact_name = data.get("contact_name", "").strip()
    sender_name = data.get("sender_name", "").strip()
    missing_items = data.get("missing_items", [])
    overpriced_items = []
    for o in data.get("overpriced_items", []):
        try:
            if float(o.get("estimate_price", 0)) > float(o.get("reference_price", 0)):
                overpriced_items.append(o)
        except (TypeError, ValueError):
            pass

    if not company_name or not contact_name or not sender_name:
        return jsonify({"error": "会社名・担当者名・送信者名を入力してください。"}), 400

    missing_lines = "\n".join(
        [f"  ・{m['item']}（{m.get('reason', '')}）" for m in missing_items]
    ) if missing_items else "  なし"

    over_lines = "\n".join(
        [f"  ・{o['item']}　見積:{int(float(o['estimate_price'])):,}円/{o.get('unit','')}　基準:{int(float(o['reference_price'])):,}円"
         for o in overpriced_items]
    ) if overpriced_items else "  なし"

    prompt = f"""以下の見積精査結果をもとに、業者への訂正依頼メールを作成してください。

[宛先情報]
会社名: {company_name}
担当者名: {contact_name}
送信者名（コスモバンク）: {sender_name}

[精査結果: 見積書に追加すべき項目]
{missing_lines}

[精査結果: 単価調整を依頼したい項目]
{over_lines}

以下の過去のメール例を参考に、同じ書式・文体でメールを作成してください。

{SAMPLE_EMAILS}

書式ルール:
- 【追加】【不要】【変更】【単価調整希望】などのセクションを必要なものだけ使う
- 単価調整の形式は「項目名　現在単価（数字のみ）　⇒　希望単価（数字のみ）」
- 各項目は簡潔に（理由・説明は不要）
- 「単価に関してはあくまでも希望ですのでご承知おきください。」は単価調整がある場合のみ記載
- 署名は送信者名のみ
- 前置きや説明は不要。メール本文のみ出力してください"""

    try:
        client = anthropic.Anthropic(api_key=api_key)
        message = client.messages.create(
            model="claude-sonnet-4-6",
            max_tokens=1024,
            messages=[{"role": "user", "content": prompt}]
        )
    except Exception as e:
        return jsonify({"error": f"AI APIの呼び出しに失敗しました: {e}"}), 500

    return jsonify({"email": message.content[0].text.strip()})


@app.route("/api/prices/match", methods=["POST"])
def match_prices():
    api_key = os.environ.get("ANTHROPIC_API_KEY")
    if not api_key:
        return jsonify({"error": "ANTHROPIC_API_KEY が設定されていません。"}), 500

    data = request.get_json()
    estimate_prices = data.get("estimate_prices", [])
    if not estimate_prices:
        return jsonify({"matches": []})

    prices = [p.to_dict() for p in Price.query.order_by(Price.id).all()]
    master_json = json.dumps(prices, ensure_ascii=False, indent=2)
    estimate_json = json.dumps(estimate_prices, ensure_ascii=False, indent=2)

    prompt = f"""以下の「見積書から抽出した工事単価」と「社内単価マスタ」を照合してください。
表記ゆれ・略語・類似表現（例:「壁紙張替え」≒「クロス貼替」、「CF貼り替え」≒「CF貼替」）も同一工事と見なして照合してください。

【社内単価マスタ】
{master_json}

【見積書から抽出した工事単価】
{estimate_json}

各見積書項目について以下のルールでactionを判定してください：
- update: マスタに同一工事が存在し、かつ見積単価がマスタ単価より安い場合（最安値更新候補）
- add: マスタに対応する工事が存在しない場合（新規追加候補）
- skip: マスタに同一工事が存在するが、見積単価がマスタ単価以上の場合（取り込み不要）

addの場合、マスタに追加する際に使う推奨項目名（部屋番号・場所の接頭語を除いたシンプルな工事名）をsuggested_itemに入れてください。
updateとskipの場合はsuggested_itemにnullを入れてください。

以下のJSON形式のみで返してください。説明文は不要です。

{{
  "matches": [
    {{
      "estimate_item": "見積書上の項目名",
      "estimate_price": 単価数値,
      "unit": "単位",
      "category": "カテゴリ名",
      "master_id": マスタIDまたはnull,
      "master_item": "マスタ上の項目名またはnull",
      "current_price": 現在のマスタ単価またはnull,
      "action": "update" または "add" または "skip",
      "suggested_item": "addの場合の推奨マスタ登録名。それ以外はnull"
    }}
  ]
}}"""

    try:
        client = anthropic.Anthropic(api_key=api_key)
        message = client.messages.create(
            model="claude-sonnet-4-6",
            max_tokens=8192,
            messages=[{"role": "user", "content": prompt}]
        )
    except Exception as e:
        return jsonify({"error": f"AI APIの呼び出しに失敗しました: {e}"}), 500

    raw = message.content[0].text.strip()
    json_match = re.search(r'\{.*\}', raw, re.DOTALL)
    if json_match:
        raw = json_match.group(0)

    try:
        result = json.loads(raw)
    except json.JSONDecodeError:
        return jsonify({"error": "照合結果の解析に失敗しました。もう一度お試しください。"}), 500

    return jsonify(result)


@app.route("/api/prices/bulk-apply", methods=["POST"])
def bulk_apply_prices():
    data = request.get_json()
    items = data.get("items", [])
    updated = 0
    added = 0
    next_id = (db.session.query(db.func.max(Price.id)).scalar() or 0) + 1

    for item in items:
        if item["action"] == "update":
            p = Price.query.get(item["master_id"])
            if p:
                p.price = int(item["estimate_price"])
                updated += 1
        elif item["action"] == "add":
            item_name = item.get("suggested_item") or item["estimate_item"]
            db.session.add(Price(
                id=next_id,
                category=item.get("category", "その他"),
                item=item_name,
                unit=item.get("unit", "式"),
                price=int(item["estimate_price"]),
                memo=""
            ))
            next_id += 1
            added += 1

    db.session.commit()
    return jsonify({"updated": updated, "added": added})


@app.route("/api/prices/extract-from-estimate", methods=["POST"])
def extract_prices_from_estimate():
    api_key = os.environ.get("ANTHROPIC_API_KEY")
    if not api_key:
        return jsonify({"error": "ANTHROPIC_API_KEY が設定されていません。"}), 500

    f = request.files.get("file")
    if not f:
        return jsonify({"error": "ファイルがありません"}), 400

    file_bytes = f.read()
    filename = f.filename.lower()

    try:
        if filename.endswith((".xlsx", ".xls")):
            text = extract_excel_text(file_bytes)
        elif filename.endswith(".pdf"):
            text = extract_pdf_text(file_bytes)
        else:
            return jsonify({"error": "対応していないファイル形式です（.xlsx/.xls/.pdf のみ）"}), 400
    except Exception as e:
        return jsonify({"error": f"ファイルの読み込みに失敗しました: {e}"}), 400

    prices = [p.to_dict() for p in Price.query.order_by(Price.id).all()]
    price_master_text = build_price_master_text(prices)

    prompt = f"""以下の業者見積書から、単価が明記されている工事項目を抽出してください。
合計金額のみで単価が不明な項目は含めないでください。

{price_master_text}

見積書内容:
{text}

抽出ルール:
1. 単価が明記されている工事項目のみ抽出
2. 項目名から部屋番号・場所の接頭語を除去してシンプルな工事名に正規化する
   例: 「和室①クロス張替え」「和室②クロス張替え」→「クロス貼替」（1件）
   例: 「洋室CFシート張替」「キッチンCF張替」→「CF貼替」（1件）
   例: 「浴室コーキング打替」→「コーキング打替」
3. 同じ工事内容が複数行ある場合（部屋違い・施工箇所違い等）は1件にまとめ、最低単価を採用する
4. 社内単価マスタに同じ工事がある場合は、マスタの表記に合わせた項目名を優先する

以下のJSON形式のみで返してください。説明文は不要です。

{{
  "estimate_prices": [
    {{"category": "カテゴリ名", "item": "正規化した工事項目名", "unit": "単位", "price": 単価数値}}
  ],
  "contractor_name": "見積書の発行会社名。見当たらない場合はnull"
}}

カテゴリは工事内容から推定してください（クロス/床/塗装/設備/建具/電気/水道/クリーニング/その他 など）。
単価が範囲で書かれている場合は中間値を使用してください。
contractor_nameは見積書のヘッダー・社名欄・発行者情報などから業者の会社名を抽出してください。見つからない場合はnull。"""

    try:
        client = anthropic.Anthropic(api_key=api_key)
        message = client.messages.create(
            model="claude-sonnet-4-6",
            max_tokens=4096,
            messages=[{"role": "user", "content": prompt}]
        )
    except Exception as e:
        return jsonify({"error": f"AI APIの呼び出しに失敗しました: {e}"}), 500

    raw = message.content[0].text.strip()
    json_match = re.search(r'\{.*\}', raw, re.DOTALL)
    if json_match:
        raw = json_match.group(0)

    try:
        result = json.loads(raw)
    except json.JSONDecodeError:
        return jsonify({"error": "AIの応答を解析できませんでした。もう一度お試しください。"}), 500

    contractor_name = result.get("contractor_name")
    if contractor_name:
        cats = list({p["category"] for p in result.get("estimate_prices", []) if p.get("category")})
        auto_register_contractor(contractor_name, cats)

    return jsonify(result)


@app.route("/api/prices/import", methods=["POST"])
def import_prices():
    f = request.files.get("file")
    if not f:
        return jsonify({"error": "ファイルがありません"}), 400

    file_bytes = f.read()
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    ws = wb.active

    next_id = (db.session.query(db.func.max(Price.id)).scalar() or 0) + 1
    added = 0

    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[0]:
            continue
        category = str(row[0]).strip() if row[0] else ""
        item = str(row[1]).strip() if len(row) > 1 and row[1] else ""
        unit = str(row[2]).strip() if len(row) > 2 and row[2] else "式"
        price_val = row[3] if len(row) > 3 and row[3] else 0
        try:
            price_int = int(float(str(price_val).replace(",", "")))
        except (ValueError, TypeError):
            continue
        if category and item and price_int > 0:
            db.session.add(Price(id=next_id, category=category, item=item, unit=unit, price=price_int, memo=""))
            next_id += 1
            added += 1

    db.session.commit()
    return jsonify({"added": added})


@app.route("/history")
def history_page():
    histories = AnalysisHistory.query.order_by(AnalysisHistory.id.desc()).limit(100).all()
    return render_template("history.html", histories=histories)


@app.route("/api/monday/status", methods=["GET"])
def monday_status():
    enabled = bool(os.environ.get("MONDAY_API_TOKEN") and os.environ.get("MONDAY_BOARD_ID"))
    return jsonify({"enabled": enabled})


@app.route("/api/monday/create-item", methods=["POST"])
def monday_create_item():
    import urllib.request as urlreq
    token = os.environ.get("MONDAY_API_TOKEN")
    board_id = os.environ.get("MONDAY_BOARD_ID")
    if not token or not board_id:
        return jsonify({"error": "MONDAY_API_TOKEN または MONDAY_BOARD_ID が設定されていません。"}), 400

    data = request.get_json()
    contractor = data.get("contractor_name") or "不明"
    file_names = data.get("file_names", "")
    missing_count = data.get("missing_count", 0)
    overpriced_count = data.get("overpriced_count", 0)
    summary = data.get("summary", "")
    analyzed_at = data.get("analyzed_at", datetime.date.today().isoformat())

    item_name = f"{contractor} ({str(analyzed_at)[:10]})"
    create_payload = json.dumps({
        "query": f"mutation {{ create_item (board_id: {board_id}, item_name: {json.dumps(item_name)}) {{ id }} }}"
    }).encode("utf-8")

    req = urlreq.Request(
        "https://api.monday.com/v2",
        data=create_payload,
        headers={"Authorization": token, "Content-Type": "application/json"}
    )
    try:
        with urlreq.urlopen(req, timeout=10) as resp:
            api_result = json.loads(resp.read())
    except Exception as e:
        return jsonify({"error": f"Monday.com API エラー: {e}"}), 500

    item_id = api_result.get("data", {}).get("create_item", {}).get("id")
    if not item_id:
        errors = api_result.get("errors", [])
        return jsonify({"error": f"アイテム作成に失敗しました: {errors}"}), 500

    detail = f"ファイル: {file_names}\n不備・見落とし: {missing_count}件\n割高項目: {overpriced_count}件\n\n{summary}"
    update_payload = json.dumps({
        "query": f"mutation {{ create_update (item_id: {item_id}, body: {json.dumps(detail)}) {{ id }} }}"
    }).encode("utf-8")
    update_req = urlreq.Request(
        "https://api.monday.com/v2",
        data=update_payload,
        headers={"Authorization": token, "Content-Type": "application/json"}
    )
    try:
        with urlreq.urlopen(update_req, timeout=10):
            pass
    except Exception:
        pass

    return jsonify({"ok": True, "item_id": item_id})


# --- 業者マスタ API ---

@app.route("/contractors")
def contractors_page():
    contractors = [c.to_dict() for c in Contractor.query.order_by(Contractor.id).all()]
    from sqlalchemy import func
    stats_rows = db.session.query(
        AnalysisHistory.contractor_name,
        func.count(AnalysisHistory.id).label("count"),
        func.avg(AnalysisHistory.missing_count).label("avg_missing"),
        func.avg(AnalysisHistory.overpriced_count).label("avg_overpriced")
    ).group_by(AnalysisHistory.contractor_name).all()
    stats_map = {
        row.contractor_name: {
            "count": row.count,
            "avg_missing": round(row.avg_missing, 1) if row.avg_missing else 0,
            "avg_overpriced": round(row.avg_overpriced, 1) if row.avg_overpriced else 0
        }
        for row in stats_rows
    }
    return render_template("contractors.html", contractors=contractors, stats_map=stats_map)


@app.route("/api/contractors", methods=["GET"])
def get_contractors():
    return jsonify([c.to_dict() for c in Contractor.query.order_by(Contractor.id).all()])


@app.route("/api/contractors", methods=["POST"])
def add_contractor():
    data = request.get_json()
    name = data.get("name", "").strip()
    if not name:
        return jsonify({"error": "会社名は必須です"}), 400
    raw_cats = data.get("categories", "")
    categories = [c.strip() for c in raw_cats.split(",") if c.strip()] if raw_cats else []
    new_id = (db.session.query(db.func.max(Contractor.id)).scalar() or 0) + 1
    entry = Contractor(
        id=new_id,
        name=name,
        categories=categories,
        estimate_count=0,
        last_estimate_date="",
        rating=int(data.get("rating", 0)),
        memo=data.get("memo", "").strip()
    )
    db.session.add(entry)
    db.session.commit()
    return jsonify(entry.to_dict()), 201


@app.route("/api/contractors/<int:contractor_id>", methods=["PUT"])
def update_contractor(contractor_id):
    c = Contractor.query.get(contractor_id)
    if not c:
        return jsonify({"error": "見つかりません"}), 404
    data = request.get_json()
    c.name = data.get("name", c.name).strip()
    raw_cats = data.get("categories", "")
    if raw_cats:
        c.categories = [x.strip() for x in raw_cats.split(",") if x.strip()]
    c.rating = int(data.get("rating", c.rating))
    c.memo = data.get("memo", c.memo or "").strip()
    db.session.commit()
    return jsonify(c.to_dict())


@app.route("/api/contractors/<int:contractor_id>", methods=["DELETE"])
def delete_contractor(contractor_id):
    Contractor.query.filter_by(id=contractor_id).delete()
    db.session.commit()
    return jsonify({"ok": True})


if __name__ == "__main__":
    app.run(host="0.0.0.0", debug=os.environ.get("FLASK_DEBUG", "false").lower() == "true", port=8080)
